"""Render pricing impact reports and RI forecasts as Excel workbooks."""

from __future__ import annotations

from pathlib import Path
from numbers import Real
from typing import Any

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.pricesheet import PricesheetMetadata
from core.usage_data import UsageMetadata
from pricing_updates.price_impact.impact import (
    RESERVATION_INVENTORY_COLUMNS,
    SUMMARY_IMPACT_COLUMNS,
    build_impact_summaries,
    get_single_pricing_currency,
)
from pricing_updates.price_impact.ri_evolution import (
    PRICE_EVOLUTION_FALLBACKS,
    PriceEvolutionResult,
)

SUMMARY_SHEET_NAME = "Non-RI Pricing Impact"
RI_PRICE_EVOLUTION_SHEET_NAME = "RI Price Evolution"
STORAGE_TABLE_NAME = "StorageData"
VM_TABLE_NAME = "VirtualMachinesData"
PRICING_MODEL_FILTER_NAME = "SelectedPricingModel"
SUMMARY_VALUE_COLUMN_WIDTH = 16
SUMMARY_NOTE_ROWS = 6
DISCLAIMER_TEXT = (
    "Disclaimer: Independent tool, not a Microsoft service or product; "
    "provided \"as is,\" with no warranties. Figures are indicative estimates "
    "based on past observed usage and may contain mapping or calculation "
    "inaccuracies. Actual charges may differ. Validate before using for "
    "budgeting or commercial decisions."
)

# One (data table, VM Impact value) source per SUMMARY_IMPACT_COLUMNS entry.
SUMMARY_IMPACT_SOURCES = (
    (STORAGE_TABLE_NAME, None),
    (VM_TABLE_NAME, "Region Price Update"),
    (VM_TABLE_NAME, "v1-v2 Price Update"),
)

PRICE_EVOLUTION_DISPLAY_MODES = ("Impact only", "Impact + baseline")


PRICE_EVOLUTION_CHART_CATEGORIES = (
    "Baseline",
    "RI expiration to RI 3Y",
    "RI expiration to SP 3Y",
    "RI expiration to PAYG",
    "Retired VM - pricing assumed to continue",
    "Impact - retired VM pricing assumption",
)


PRICE_EVOLUTION_CHART_COLORS = (
    "4472C4",
    "70AD47",
    "8064A2",
    "4BACC6",
    "F4B183",
    "5B9BD5",
)


def _summary_cell_value(value: Any) -> Any:
    if value is None or value is pd.NA or pd.isna(value):
        return None
    if isinstance(value, Real) and not isinstance(value, bool):
        return float(value)
    return value


def _write_summary_block(
    worksheet: Any,
    row: int,
    column: int,
    title: str,
    headers: list[str],
    rows: list[list[Any]],
) -> int:
    """Write a titled block and return the row following its last row."""
    worksheet.cell(row=row, column=column, value=title).font = Font(
        bold=True,
        size=12,
    )
    for offset, header in enumerate(headers):
        worksheet.cell(
            row=row + 1,
            column=column + offset,
            value=header,
        ).font = Font(bold=True)
    for row_offset, values in enumerate(rows, start=2):
        for offset, value in enumerate(values):
            worksheet.cell(
                row=row + row_offset,
                column=column + offset,
                value=_summary_cell_value(value),
            )
    return row + 2 + len(rows)


def _pricesheet_metadata_rows(
    metadata: PricesheetMetadata,
) -> list[list[Any]]:
    def date_text(value: Any) -> str:
        return (
            "N/A"
            if value is None or pd.isna(value)
            else pd.Timestamp(value).date().isoformat()
        )

    source_files = [Path(path).name for path in metadata.source_files or []]
    shown_files = "; ".join(source_files[:3])
    if len(source_files) > 3:
        shown_files += f" (+{len(source_files) - 3} more)"
    return [
        [
            "Contract Type",
            metadata.contract_type.upper() if metadata.contract_type else "N/A",
        ],
        ["Effective Start", date_text(metadata.effective_start_date)],
        ["Effective End", date_text(metadata.effective_end_date)],
        ["Source Files", shown_files or "N/A"],
    ]


def _impact_scope_note(ri_price_evolution_enabled: bool) -> str:
    if ri_price_evolution_enabled:
        reservation_hint = (
            f"To see the impact of reservations, open the "
            f"'{RI_PRICE_EVOLUTION_SHEET_NAME}' tab."
        )
    else:
        reservation_hint = (
            "This file was generated without the RI price evolution option. "
            "To see the impact of reservations, rerun the script with "
            "--ri-price-evolution (requires --pricesheet and "
            "--reservation-inventory)."
        )
    return (
        "This is the initial expected impact effective 1 February 2027 for "
        "non-reservation pricing models only. Reserved Instance costs are "
        "excluded because existing RIs keep their current price until "
        "expiration; their impact will increase as RIs expire. "
        + reservation_hint
        + " Savings Plan figures follow the PAYG price increase, which assumes "
        "additional Savings Plan commitment is purchased to keep the same "
        "workload covered; keeping only the existing commitment would cost more."
    )


def _summary_sumifs(
    table: str,
    available_columns: set[str],
    group_column: str,
    label_cell: str,
    impact_type: str | None,
) -> str:
    """Build a SUMIFS over a data table, or 0 if the table cannot be filtered."""
    required = {group_column, "PricingModel"}
    if impact_type is not None:
        required.add("Impact")
    if not required.issubset(available_columns):
        return "0"
    criteria = [f"{table}[{group_column}],{label_cell}"]
    if impact_type is not None:
        criteria.append(f'{table}[Impact],"{impact_type}"')
    criteria.append(
        f"{table}[PricingModel],"
        f'IF({PRICING_MODEL_FILTER_NAME}="All","<>Reservation",'
        f"{PRICING_MODEL_FILTER_NAME})"
    )
    return f"SUMIFS({table}[CostImpact],{','.join(criteria)})"


def _write_formula_section(
    worksheet: Any,
    row: int,
    title: str,
    frame: pd.DataFrame,
    group_column: str,
    available_columns: dict[str, set[str]],
) -> int:
    """Write an impact section whose values are formulas over the data tables."""
    value_count = len([*SUMMARY_IMPACT_COLUMNS, "Total Impact"])
    first_value_column = len(frame.columns) - value_count + 1
    label_column_letter = get_column_letter(
        list(frame.columns).index(group_column) + 1
    )
    worksheet.cell(row=row, column=1, value=title).font = Font(
        bold=True,
        size=12,
    )
    for index, header in enumerate(frame.columns, start=1):
        worksheet.cell(row=row + 1, column=index, value=header).font = Font(
            bold=True,
        )
    first_data_row = row + 2
    data_frame = frame.iloc[:-1]
    for offset, values in enumerate(data_frame.itertuples(index=False)):
        current_row = first_data_row + offset
        for index, value in enumerate(values, start=1):
            if index < first_value_column:
                worksheet.cell(
                    row=current_row,
                    column=index,
                    value=_summary_cell_value(value),
                )
        label_cell = f"${label_column_letter}{current_row}"
        for index, (table, impact_type) in enumerate(
            SUMMARY_IMPACT_SOURCES,
            start=first_value_column,
        ):
            worksheet.cell(
                row=current_row,
                column=index,
                value="="
                + _summary_sumifs(
                    table,
                    available_columns[table],
                    group_column,
                    label_cell,
                    impact_type,
                ),
            )
        total_column = first_value_column + len(SUMMARY_IMPACT_COLUMNS)
        worksheet.cell(
            row=current_row,
            column=total_column,
            value=(
                f"=SUM({get_column_letter(first_value_column)}{current_row}:"
                f"{get_column_letter(total_column - 1)}{current_row})"
            ),
        )
    total_row = first_data_row + len(data_frame)
    for index, value in enumerate(frame.iloc[-1], start=1):
        if index < first_value_column:
            worksheet.cell(
                row=total_row,
                column=index,
                value=_summary_cell_value(value),
            )
    for index in range(first_value_column, first_value_column + value_count):
        letter = get_column_letter(index)
        worksheet.cell(
            row=total_row,
            column=index,
            value=(
                f"=SUM({letter}{first_data_row}:{letter}{total_row - 1})"
                if len(data_frame)
                else 0
            ),
        )
    for cell in worksheet[total_row]:
        cell.font = Font(bold=True)
    for cells in worksheet.iter_rows(
        min_row=first_data_row,
        max_row=total_row,
        min_col=first_value_column,
        max_col=first_value_column + value_count - 1,
    ):
        for cell in cells:
            cell.number_format = "#,##0.00"
    return total_row + 3


def _add_data_table(
    worksheet: Any,
    table_name: str,
    dataframe: pd.DataFrame,
) -> None:
    """Expose a written dataframe sheet as an Excel Table."""
    if len(dataframe.columns) == 0:
        return
    # Excel tables need at least one data row; empty frames get a blank one.
    last_row = max(2, len(dataframe) + 1)
    table = Table(
        displayName=table_name,
        ref=f"A1:{get_column_letter(len(dataframe.columns))}{last_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
    worksheet.freeze_panes = "A2"


def _write_summary_sheet(
    writer: pd.ExcelWriter,
    storage_report: pd.DataFrame,
    vm_report: pd.DataFrame,
    usage_metadata: UsageMetadata,
    pricing_currency: str,
    price_protection_enabled: bool,
    pricesheet_metadata: PricesheetMetadata | None = None,
    ri_price_evolution_enabled: bool = False,
) -> None:
    required_metadata = {
        "contract_type": usage_metadata.contract_type,
        "usage_start_date": usage_metadata.usage_start_date,
        "usage_end_date": usage_metadata.usage_end_date,
    }
    missing_metadata = [
        name
        for name, value in required_metadata.items()
        if value is None
    ]
    if missing_metadata:
        raise ValueError(
            "Usage metadata is missing required fields: "
            + ", ".join(missing_metadata)
        )
    metadata_card = pd.DataFrame(
        {
            "Field": [
                "Pricing Currency",
                "Contract Type",
                "Price Protection",
                "Usage Start",
                "Usage End",
            ],
            "Value": [
                pricing_currency,
                usage_metadata.contract_type.upper(),
                "Enabled" if price_protection_enabled else "Disabled",
                pd.Timestamp(usage_metadata.usage_start_date)
                .date()
                .isoformat(),
                pd.Timestamp(usage_metadata.usage_end_date)
                .date()
                .isoformat(),
            ],
        }
    )
    overall, business, subscriptions = build_impact_summaries(
        storage_report,
        vm_report,
    )
    worksheet = writer.book.create_sheet(SUMMARY_SHEET_NAME)
    writer.sheets[SUMMARY_SHEET_NAME] = worksheet

    next_row = max(
        _write_summary_block(
            worksheet,
            1,
            1,
            "Usage Metadata",
            list(metadata_card.columns),
            metadata_card.values.tolist(),
        ),
        (
            _write_summary_block(
                worksheet,
                1,
                4,
                "Pricesheet Metadata",
                ["Field", "Value"],
                _pricesheet_metadata_rows(pricesheet_metadata),
            )
            if pricesheet_metadata is not None
            else 0
        ),
    )

    note_title_row = next_row + 1
    worksheet.cell(
        row=note_title_row,
        column=1,
        value="Impact Scope",
    ).font = Font(bold=True, size=12)
    note_row = note_title_row + 1
    worksheet.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row + SUMMARY_NOTE_ROWS - 1,
        end_column=6,
    )
    note_cell = worksheet.cell(
        row=note_row,
        column=1,
        value=_impact_scope_note(ri_price_evolution_enabled),
    )
    note_cell.alignment = Alignment(wrap_text=True, vertical="top")

    disclaimer_row = note_row + SUMMARY_NOTE_ROWS + 1
    _write_disclaimer(worksheet, disclaimer_row, 6, rows=3)

    overall_title_row = disclaimer_row + 4
    overall_rows = overall.iloc[:-1]
    pricing_models = [
        str(model) for model in overall_rows["Pricing Model"].dropna()
    ]
    after_overall = _write_summary_block(
        worksheet,
        overall_title_row,
        1,
        "Overall Impact",
        list(overall.columns),
        overall.values.tolist(),
    )
    overall_total_row = after_overall - 1
    for cells in worksheet.iter_rows(
        min_row=overall_title_row + 2,
        max_row=overall_total_row,
        min_col=2,
        max_col=len(overall.columns),
    ):
        for cell in cells:
            cell.number_format = "#,##0.00"
    for cell in worksheet[overall_total_row]:
        cell.font = Font(bold=True)

    selector_row = after_overall + 1
    worksheet.cell(
        row=selector_row,
        column=1,
        value="Pricing model filter",
    ).font = Font(bold=True)
    selector = worksheet.cell(row=selector_row, column=2, value="All")
    selector.fill = PatternFill("solid", fgColor="FFF2CC")
    worksheet.cell(
        row=selector_row,
        column=3,
        value=(
            "Applies to Impact by Business Entity and Impact by "
            "Subscription. All excludes Reservation."
        ),
    ).font = Font(italic=True)
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(["All", *pricing_models]) + '"',
        allow_blank=False,
    )
    validation.error = "Select a value from the list."
    validation.errorTitle = "Invalid selection"
    validation.showErrorMessage = True
    worksheet.add_data_validation(validation)
    validation.add(selector)
    writer.book.defined_names.add(
        DefinedName(
            PRICING_MODEL_FILTER_NAME,
            attr_text=(
                f"'{SUMMARY_SHEET_NAME}'!$B${selector_row}"
            ),
        )
    )

    available_columns = {
        STORAGE_TABLE_NAME: set(storage_report.columns),
        VM_TABLE_NAME: set(vm_report.columns),
    }
    next_row = selector_row + 2
    if business is not None:
        next_row = _write_formula_section(
            worksheet,
            next_row,
            "Impact by Business Entity",
            business,
            "BusinessEntity",
            available_columns,
        )
    _write_formula_section(
        worksheet,
        next_row,
        "Impact by Subscription",
        subscriptions,
        "SubscriptionId",
        available_columns,
    )

    for column_cells in worksheet.columns:
        width = min(
            max(
                len(str(cell.value))
                if cell.value is not None
                and cell.row not in (note_row, selector_row, disclaimer_row)
                else 0
                for cell in column_cells
            )
            + 2,
            45,
        )
        worksheet.column_dimensions[column_cells[0].column_letter].width = width
    for column in range(2, 8):
        worksheet.column_dimensions[get_column_letter(column)].width = (
            SUMMARY_VALUE_COLUMN_WIDTH
        )
    # Narrow columns: wrap the bold header labels instead of widening them.
    for cells in worksheet.iter_rows(min_row=2, min_col=2, max_col=7):
        for cell in cells:
            if cell.font.bold and isinstance(cell.value, str):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in range(note_row, note_row + SUMMARY_NOTE_ROWS):
        worksheet.row_dimensions[row].height = 18
    for row in range(disclaimer_row, disclaimer_row + 3):
        worksheet.row_dimensions[row].height = 18


def _write_disclaimer(
    worksheet: Any,
    row: int,
    end_column: int,
    rows: int = 1,
) -> None:
    """Write the merged, wrapped disclaimer starting at column A."""
    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row + rows - 1,
        end_column=end_column,
    )
    cell = worksheet.cell(row=row, column=1, value=DISCLAIMER_TEXT)
    cell.font = Font(italic=True)
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def _add_price_evolution_table(
    worksheet: Any,
    table_name: str,
) -> None:
    if worksheet.max_row < 2 or worksheet.max_column < 1:
        return
    table = Table(displayName=table_name, ref=worksheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
    worksheet.freeze_panes = "A2"


def _price_evolution_column_ranges(
    dataframe: pd.DataFrame,
    sheet_name: str,
) -> dict[str, str]:
    last_row = max(2, len(dataframe) + 1)
    escaped_sheet = sheet_name.replace("'", "''")
    return {
        column: (
            f"'{escaped_sheet}'!"
            f"${get_column_letter(index)}$2:"
            f"${get_column_letter(index)}${last_row}"
        )
        for index, column in enumerate(dataframe.columns, start=1)
    }


def _add_price_evolution_selector(
    workbook: Any,
    dashboard: Any,
    lists_sheet: Any,
    *,
    row: int,
    list_column: int,
    label: str,
    name: str,
    values: list[str],
    default: str,
) -> None:
    dashboard.cell(row=row, column=1, value=label).font = Font(bold=True)
    dashboard.cell(row=row, column=2, value=default)
    lists_sheet.cell(row=1, column=list_column, value=label)
    for value_row, value in enumerate(values, start=2):
        lists_sheet.cell(
            row=value_row,
            column=list_column,
            value=value,
        )
    column_letter = get_column_letter(list_column)
    defined_name = DefinedName(
        name,
        attr_text=(
            f"'RI Evolution Lists'!"
            f"${column_letter}$2:${column_letter}${len(values) + 1}"
        ),
    )
    workbook.defined_names.add(defined_name)
    validation = DataValidation(
        type="list",
        formula1=f"={name}",
        allow_blank=False,
    )
    validation.error = "Select a value from the list."
    validation.errorTitle = "Invalid selection"
    validation.showErrorMessage = True
    dashboard.add_data_validation(validation)
    validation.add(dashboard.cell(row=row, column=2))


def _price_evolution_sumifs_formula(
    sum_range: str,
    criteria: list[tuple[str, str]],
) -> str:
    arguments = [sum_range]
    for criteria_range, criterion in criteria:
        arguments.extend([criteria_range, criterion])
    return "=SUMIFS(" + ",".join(arguments) + ")"


def _write_price_evolution_dashboard(
    writer: pd.ExcelWriter,
    result: PriceEvolutionResult,
    pricing_currency: str,
) -> None:
    workbook = writer.book
    dashboard = workbook.create_sheet(RI_PRICE_EVOLUTION_SHEET_NAME, 1)
    writer.sheets[RI_PRICE_EVOLUTION_SHEET_NAME] = dashboard
    lists_sheet = workbook.create_sheet("RI Evolution Lists")
    writer.sheets["RI Evolution Lists"] = lists_sheet

    dashboard["A1"] = "RI Price Evolution - Keep Until Expiration"
    dashboard["A1"].font = Font(bold=True, size=16)
    dashboard["D1"] = "Forecast"
    dashboard["D1"].font = Font(bold=True)
    dashboard["E1"] = (
        f"{result.forecast_start:%Y-%m} to "
        f"{result.forecast_end:%Y-%m}"
    )
    dashboard["D2"] = "Currency"
    dashboard["E2"] = pricing_currency
    dashboard["D3"] = "Fiscal year starts"
    dashboard["E3"] = pd.Timestamp(
        year=2000,
        month=result.fiscal_year_start_month,
        day=1,
    ).strftime("%B")
    dashboard["D4"] = "Lifecycle source"
    dashboard["E4"] = result.lifecycle_source
    dashboard["D5"] = "Lifecycle snapshot"
    dashboard["E5"] = result.lifecycle_snapshot_date

    selector_definitions = [
        (
            "FallbackPolicy",
            "Fallback after RI expiration",
            "RIPE_Fallback",
            PRICE_EVOLUTION_FALLBACKS,
            "RI 3Y",
        ),
        (
            "DisplayMode",
            "Display mode",
            "RIPE_DisplayMode",
            PRICE_EVOLUTION_DISPLAY_MODES,
            "Impact only",
        ),
        (
            "ImpactClassification",
            "Price impact",
            "RIPE_Impact",
            None,
            "All",
        ),
        (
            "BusinessEntity",
            "Business entity",
            "RIPE_BusinessEntity",
            None,
            "All",
        ),
        (
            "SubscriptionName",
            "Subscription",
            "RIPE_Subscription",
            None,
            "All",
        ),
        (
            "Region",
            "Region",
            "RIPE_Region",
            None,
            "All",
        ),
        (
            "VMFamily",
            "VM family",
            "RIPE_VMFamily",
            None,
            "All",
        ),
        (
            "VMSize",
            "VM size",
            "RIPE_VMSize",
            None,
            "All",
        ),
        (
            "FiscalYear",
            "Fiscal year",
            "RIPE_FiscalYear",
            None,
            "All",
        ),
    ]
    chart_ranges = _price_evolution_column_ranges(
        result.chart_data,
        "RI Evolution Chart Data",
    )
    projection_ranges = _price_evolution_column_ranges(
        result.projection,
        "RI Projection Details",
    )
    selector_criteria: list[tuple[str, str]] = []
    projection_criteria: list[tuple[str, str]] = []
    for offset, (column, label, name, options, default) in enumerate(
        selector_definitions,
        start=2,
    ):
        values = (
            list(options)
            if options is not None
            else [
                "All",
                *sorted(
                    result.projection[column].dropna().astype(str).unique(),
                    key=None if column == "FiscalYear" else str.casefold,
                ),
            ]
        )
        _add_price_evolution_selector(
            workbook,
            dashboard,
            lists_sheet,
            row=offset,
            list_column=offset - 1,
            label=label,
            name=name,
            values=list(dict.fromkeys(values)),
            default=default,
        )
        cell = f"$B${offset}"
        criterion = f'IF({cell}="All","*",{cell})' if options is None else cell
        selector_criteria.append((chart_ranges[column], criterion))
        if column != "DisplayMode":
            projection_criteria.append((projection_ranges[column], criterion))

    dashboard["A12"] = (
        "VM RI baseline only (used + unused). Unpriced cohorts are excluded "
        "from chart values; totals are marked Unpriced."
    )
    _write_disclaimer(dashboard, 13, 9)
    dashboard.row_dimensions[13].height = 32
    table_start_row = 14
    dashboard.cell(
        row=table_start_row,
        column=1,
        value="Forecast month",
    )
    for column_offset, category in enumerate(
        PRICE_EVOLUTION_CHART_CATEGORIES,
        start=2,
    ):
        dashboard.cell(
            row=table_start_row,
            column=column_offset,
            value=category,
        )
    total_column = len(PRICE_EVOLUTION_CHART_CATEGORIES) + 2
    dashboard.cell(
        row=table_start_row,
        column=total_column,
        value="Total",
    )
    unpriced_column = total_column + 1
    dashboard.cell(
        row=table_start_row,
        column=unpriced_column,
        value="Unpriced cohorts",
    )
    for cell in dashboard[table_start_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")

    for month_offset, month in enumerate(
        pd.date_range(
            result.forecast_start,
            result.forecast_end,
            freq="MS",
        ),
        start=table_start_row + 1,
    ):
        dashboard.cell(row=month_offset, column=1, value=month.to_pydatetime())
        dashboard.cell(
            row=month_offset,
            column=1,
        ).number_format = "mmm-yy"
        for category_offset, category in enumerate(
            PRICE_EVOLUTION_CHART_CATEGORIES,
            start=2,
        ):
            formula = _price_evolution_sumifs_formula(
                chart_ranges["DisplayAmount"],
                [
                    (
                        chart_ranges["ForecastMonth"],
                        f"$A{month_offset}",
                    ),
                    (
                        chart_ranges["DisplayCategory"],
                        f"{get_column_letter(category_offset)}"
                        f"${table_start_row}",
                    ),
                    *selector_criteria,
                ],
            )
            dashboard.cell(
                row=month_offset,
                column=category_offset,
                value=formula,
            ).number_format = "#,##0.00"
        first_value = get_column_letter(2)
        last_value = get_column_letter(total_column - 1)
        dashboard.cell(
            row=month_offset,
            column=unpriced_column,
            value=_price_evolution_sumifs_formula(
                projection_ranges["UnpricedCohortCount"],
                [
                    (projection_ranges["ForecastMonth"], f"$A{month_offset}"),
                    *projection_criteria,
                ],
            ),
        )
        unpriced_cell = f"{get_column_letter(unpriced_column)}{month_offset}"
        dashboard.cell(
            row=month_offset,
            column=total_column,
            value=f'=IF({unpriced_cell}>0,"Unpriced",'
            f"SUM({first_value}{month_offset}:{last_value}{month_offset}))",
        ).number_format = "#,##0.00"

    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.title = "Monthly RI price evolution"
    chart.y_axis.title = pricing_currency
    chart.x_axis.title = "Forecast month"
    chart.x_axis.axPos = "b"
    chart.x_axis.delete = False
    chart.x_axis.tickLblPos = "low"
    chart.x_axis.tickLblSkip = 1
    chart.x_axis.numFmt = "mmm-yy"
    chart.x_axis.numFmt.sourceLinked = False
    chart.y_axis.axPos = "l"
    chart.y_axis.delete = False
    chart.y_axis.tickLblPos = "low"
    chart.y_axis.numFmt = "#,##0.00"
    chart.y_axis.numFmt.sourceLinked = False
    chart.height = 15
    chart.width = 24
    chart.legend.position = "b"
    chart.legend.overlay = False
    last_dashboard_row = table_start_row + result.horizon_months
    data = Reference(
        dashboard,
        min_col=2,
        max_col=total_column - 1,
        min_row=table_start_row,
        max_row=last_dashboard_row,
    )
    categories = Reference(
        dashboard,
        min_col=1,
        min_row=table_start_row + 1,
        max_row=last_dashboard_row,
    )
    chart.add_data(data, titles_from_data=True)
    for series, color in zip(
        chart.series,
        PRICE_EVOLUTION_CHART_COLORS,
    ):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color
        series.invertIfNegative = False
    chart.set_categories(categories)
    dashboard.add_chart(chart, "L2")

    dashboard.freeze_panes = f"A{table_start_row + 1}"
    dashboard.column_dimensions["A"].width = 34
    dashboard.column_dimensions["B"].width = 30
    for column in range(4, 6):
        dashboard.column_dimensions[get_column_letter(column)].width = 24
    lists_sheet.sheet_state = "hidden"


def _write_price_evolution_assumptions(
    writer: pd.ExcelWriter,
    result: PriceEvolutionResult,
    usage_metadata: UsageMetadata,
) -> None:
    assumptions = pd.DataFrame(
        [
            ("Scenario", "Keep existing RIs until expiration"),
            (
                "Scope",
                "VM reservation usage and unused commitment only; existing "
                "OnDemand, Spot and Savings Plan usage is excluded",
            ),
            (
                "Projection grain",
                "Monthly totals by fallback, reporting dimensions and "
                "contribution category; no per-reservation detail",
            ),
            (
                "Unpriced cohorts",
                "Excluded from chart values; affected projected totals are "
                "marked Unpriced, not reported as complete",
            ),
            ("Forecast start", result.forecast_start.date().isoformat()),
            ("Forecast end", result.forecast_end.date().isoformat()),
            ("Horizon months", result.horizon_months),
            (
                "Fiscal year start month",
                pd.Timestamp(
                    year=2000,
                    month=result.fiscal_year_start_month,
                    day=1,
                ).strftime("%B"),
            ),
            (
                "Usage start",
                str(usage_metadata.usage_start_date),
            ),
            ("Usage end", str(usage_metadata.usage_end_date)),
            ("Price update effective date", "2027-02-01"),
            (
                "Usage assumption",
                "Aggregated input is constant reference-month usage",
            ),
            (
                "Expiration rule",
                "Expiration month applies from the first day",
            ),
            (
                "Fallback chain",
                "RI 3Y -> SP 3Y -> PAYG; SP 3Y -> PAYG",
            ),
            (
                "Unknown retirement",
                "No announced retirement; VM continues indefinitely",
            ),
            (
                "Retired PAYG",
                "Last known PAYG price continues as a high-risk assumption",
            ),
            ("Lifecycle source", result.lifecycle_source),
            ("Lifecycle snapshot", result.lifecycle_snapshot_date),
            (
                "PAYG price protection",
                (
                    "Enabled"
                    if result.price_protection_enabled
                    else "Disabled"
                ),
            ),
        ],
        columns=["Assumption", "Value"],
    )
    assumptions.to_excel(
        writer,
        sheet_name="RI Evolution Assumptions",
        index=False,
    )
    worksheet = writer.sheets["RI Evolution Assumptions"]
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    worksheet.column_dimensions["A"].width = 32
    worksheet.column_dimensions["B"].width = 80
    worksheet.freeze_panes = "A2"


def _write_price_evolution_sheets(
    writer: pd.ExcelWriter,
    result: PriceEvolutionResult,
    usage_metadata: UsageMetadata,
    pricing_currency: str,
) -> None:
    result.projection.to_excel(
        writer,
        sheet_name="RI Projection Details",
        index=False,
    )
    result.chart_data.to_excel(
        writer,
        sheet_name="RI Evolution Chart Data",
        index=False,
    )
    projection_sheet = writer.sheets["RI Projection Details"]
    chart_data_sheet = writer.sheets["RI Evolution Chart Data"]
    _add_price_evolution_table(
        projection_sheet,
        "RIPriceEvolutionProjection",
    )
    _add_price_evolution_table(
        chart_data_sheet,
        "RIPriceEvolutionChartData",
    )
    _write_price_evolution_dashboard(writer, result, pricing_currency)
    _write_price_evolution_assumptions(
        writer,
        result,
        usage_metadata,
    )
    chart_data_sheet.sheet_state = "hidden"
    writer.book.calculation.fullCalcOnLoad = True
    writer.book.calculation.forceFullCalc = True
    writer.book.calculation.calcMode = "auto"


def write_report(
    storage_report: pd.DataFrame,
    vm_report: pd.DataFrame,
    output_path: Path,
    vm_pricesheet_report: pd.DataFrame | None = None,
    storage_pricesheet_report: pd.DataFrame | None = None,
    usage_metadata: UsageMetadata | None = None,
    pricing_currency: str | None = None,
    price_protection_enabled: bool = False,
    price_evolution: PriceEvolutionResult | None = None,
    reservation_inventory: pd.DataFrame | None = None,
    pricesheet_metadata: PricesheetMetadata | None = None,
) -> None:
    """Write category reports and optional inventory and pricing worksheets."""
    if usage_metadata is None:
        raise ValueError("Usage metadata is required to write the report.")
    if pricing_currency is None:
        pricing_currency = get_single_pricing_currency(
            storage_report,
            vm_report,
        )
    storage_output = storage_report.drop(
        columns="PricingCurrency",
        errors="ignore",
    )
    vm_output = vm_report.drop(
        columns="PricingCurrency",
        errors="ignore",
    )
    reservations_output = None
    if reservation_inventory is not None:
        reservations_output = reservation_inventory.drop(
            columns=[
                *RESERVATION_INVENTORY_COLUMNS.values(),
                "Deep link to reservation",
            ],
            errors="ignore",
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _write_summary_sheet(
            writer,
            storage_output,
            vm_output,
            usage_metadata,
            pricing_currency,
            price_protection_enabled,
            pricesheet_metadata=pricesheet_metadata,
            ri_price_evolution_enabled=price_evolution is not None,
        )
        storage_output.to_excel(writer, sheet_name="Storage", index=False)
        _add_data_table(
            writer.sheets["Storage"],
            STORAGE_TABLE_NAME,
            storage_output,
        )
        vm_output.to_excel(
            writer,
            sheet_name="Virtual Machines",
            index=False,
        )
        _add_data_table(
            writer.sheets["Virtual Machines"],
            VM_TABLE_NAME,
            vm_output,
        )
        if reservations_output is not None:
            reservations_output.to_excel(
                writer,
                sheet_name="Reservations",
                index=False,
            )
        if storage_pricesheet_report is not None:
            storage_pricesheet_report.to_excel(
                writer,
                sheet_name="Storage Pricesheet",
                index=False,
            )
        if vm_pricesheet_report is not None:
            vm_pricesheet_report.to_excel(
                writer,
                sheet_name="VM Pricesheet",
                index=False,
            )
        if price_evolution is not None:
            _write_price_evolution_sheets(
                writer,
                price_evolution,
                usage_metadata,
                pricing_currency,
            )
        writer.book.calculation.fullCalcOnLoad = True
